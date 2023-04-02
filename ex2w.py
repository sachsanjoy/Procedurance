import openpyxl
#import docx
from docx import Document

# Load the Excel file
wb = openpyxl.load_workbook('book.xlsx')

# Get the active worksheet
ws = wb.active

# Create a new Word document
doc = Document()

# Loop through each row and column in the Excel file and add the data to the Word document
for row in ws.iter_rows():
    table = doc.add_table(rows=1, cols=len(row))
    hdr_cells = table.rows[0].cells
    for i in range(len(row)):
        hdr_cells[i].text = str(row[i].value)

# Save the Word document
doc.save('word_file.docx')